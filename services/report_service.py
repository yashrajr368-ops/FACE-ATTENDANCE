"""
Report Service Module
Generates professional CSV and Excel (.xlsx) attendance exports with
date range filtering, student-wise summaries, and styling.
"""

import io
import csv
from datetime import datetime, date
from typing import Optional, List, Dict
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from database import db
from models.student import Student
from models.attendance import Attendance


class ReportService:
    def get_records_for_report(
        self,
        start_date: Optional[str] = None,
        end_date: Optional[str] = None,
        student_id: Optional[str] = None,
        course: Optional[str] = None
    ) -> List[Attendance]:
        """Queries attendance records for reporting within date range and filters."""
        query = Attendance.query.join(Student, Attendance.student_id == Student.student_id)

        if start_date:
            try:
                s_date = datetime.strptime(start_date.strip(), "%Y-%m-%d").date()
                query = query.filter(Attendance.attendance_date >= s_date)
            except ValueError:
                pass

        if end_date:
            try:
                e_date = datetime.strptime(end_date.strip(), "%Y-%m-%d").date()
                query = query.filter(Attendance.attendance_date <= e_date)
            except ValueError:
                pass

        if student_id:
            query = query.filter(Student.student_id == student_id.strip())

        if course:
            query = query.filter(Student.course == course.strip())

        return query.order_by(Attendance.attendance_date.desc(), Attendance.attendance_time.desc()).all()

    def generate_csv_report(
        self,
        start_date: Optional[str] = None,
        end_date: Optional[str] = None,
        student_id: Optional[str] = None,
        course: Optional[str] = None
    ) -> io.BytesIO:
        """Generates a CSV report buffer with UTF-8 BOM encoding for seamless Excel opening."""
        records = self.get_records_for_report(start_date, end_date, student_id, course)

        output = io.StringIO()
        # Write UTF-8 BOM
        output.write('\ufeff')

        writer = csv.writer(output)
        writer.writerow([
            "Attendance ID",
            "Student ID",
            "Student Name",
            "Roll Number",
            "Course / Class",
            "Date",
            "Time",
            "Status"
        ])

        for rec in records:
            student_name = rec.student.name if rec.student else "Unknown"
            roll_no = rec.student.roll_number if rec.student else ""
            student_course = rec.student.course if rec.student else ""

            writer.writerow([
                rec.id,
                rec.student_id,
                student_name,
                roll_no,
                student_course,
                rec.attendance_date.strftime("%Y-%m-%d") if rec.attendance_date else "",
                rec.attendance_time.strftime("%H:%M:%S") if rec.attendance_time else "",
                rec.status
            ])

        bytes_io = io.BytesIO()
        bytes_io.write(output.getvalue().encode('utf-8'))
        bytes_io.seek(0)
        return bytes_io

    def generate_excel_report(
        self,
        start_date: Optional[str] = None,
        end_date: Optional[str] = None,
        student_id: Optional[str] = None,
        course: Optional[str] = None
    ) -> io.BytesIO:
        """
        Generates a beautifully styled .xlsx Excel spreadsheet with headers,
        status highlighting, and auto-adjusted column widths.
        """
        records = self.get_records_for_report(start_date, end_date, student_id, course)

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Attendance Records"

        # Styling definitions
        font_title = Font(name="Segoe UI", size=15, bold=True, color="1E293B")
        font_subtitle = Font(name="Segoe UI", size=10, italic=True, color="64748B")
        font_header = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
        font_data = Font(name="Segoe UI", size=10, color="0F172A")
        font_present = Font(name="Segoe UI", size=10, bold=True, color="15803D")

        fill_header = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
        fill_zebra = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")
        fill_present = PatternFill(start_color="DCFCE7", end_color="DCFCE7", fill_type="solid")

        align_center = Alignment(horizontal="center", vertical="center")
        align_left = Alignment(horizontal="left", vertical="center")
        thin_border = Border(
            left=Side(style='thin', color='E2E8F0'),
            right=Side(style='thin', color='E2E8F0'),
            top=Side(style='thin', color='E2E8F0'),
            bottom=Side(style='thin', color='E2E8F0')
        )

        # Title Block
        ws.merge_cells("A1:H1")
        ws["A1"] = "CAMPUS ATTENDANCE MANAGEMENT SYSTEM"
        ws["A1"].font = font_title
        ws["A1"].alignment = align_left

        ws.merge_cells("A2:H2")
        filter_desc = []
        if start_date:
            filter_desc.append(f"From: {start_date}")
        if end_date:
            filter_desc.append(f"To: {end_date}")
        if course:
            filter_desc.append(f"Course: {course}")
        if student_id:
            filter_desc.append(f"Student: {student_id}")
        filter_str = ", ".join(filter_desc) if filter_desc else "All Records"

        ws["A2"] = f"Report Generated: {datetime.now().strftime('%Y-%m-%d %H:%M')} | Filters: {filter_str} | Total Records: {len(records)}"
        ws["A2"].font = font_subtitle
        ws["A2"].alignment = align_left

        # Table Header
        headers = [
            "Record ID",
            "Student ID",
            "Student Name",
            "Roll Number",
            "Course / Class",
            "Attendance Date",
            "Attendance Time",
            "Status"
        ]

        header_row = 4
        ws.row_dimensions[header_row].height = 26
        for col_idx, header_text in enumerate(headers, 1):
            cell = ws.cell(row=header_row, column=col_idx, value=header_text)
            cell.font = font_header
            cell.fill = fill_header
            cell.alignment = align_center
            cell.border = thin_border

        # Data Rows
        current_row = 5
        for rec in records:
            ws.row_dimensions[current_row].height = 20
            student_name = rec.student.name if rec.student else "Unknown"
            roll_no = rec.student.roll_number if rec.student else ""
            student_course = rec.student.course if rec.student else ""

            row_data = [
                rec.id,
                rec.student_id,
                student_name,
                roll_no,
                student_course,
                rec.attendance_date.strftime("%Y-%m-%d") if rec.attendance_date else "",
                rec.attendance_time.strftime("%H:%M:%S") if rec.attendance_time else "",
                rec.status
            ]

            is_zebra = (current_row % 2 == 0)

            for col_idx, val in enumerate(row_data, 1):
                cell = ws.cell(row=current_row, column=col_idx, value=val)
                cell.font = font_data
                cell.border = thin_border
                if is_zebra:
                    cell.fill = fill_zebra

                if col_idx in (1, 2, 4, 6, 7):
                    cell.alignment = align_center
                elif col_idx in (3, 5):
                    cell.alignment = align_left
                elif col_idx == 8:
                    cell.alignment = align_center
                    if val == "Present":
                        cell.font = font_present
                        cell.fill = fill_present

            current_row += 1

        # Auto-adjust column widths
        for col in ws.columns:
            max_len = 0
            col_letter = get_column_letter(col[0].column)
            for cell in col:
                if cell.row >= header_row and cell.value:
                    max_len = max(max_len, len(str(cell.value)))
            ws.column_dimensions[col_letter].width = max(max_len + 4, 14)

        bytes_io = io.BytesIO()
        wb.save(bytes_io)
        bytes_io.seek(0)
        return bytes_io

    def get_student_summary_report(self) -> List[Dict]:
        """
        Generates aggregate attendance metrics per student:
        Total distinct dates with attendance recorded across campus,
        present days for this student, absent days, and percentage.
        """
        students = Student.query.order_by(Student.name.asc()).all()
        # Total distinct attendance dates in system
        total_days_res = db.session.query(db.func.count(db.distinct(Attendance.attendance_date))).scalar() or 0

        summary = []
        for s in students:
            present_days = Attendance.query.filter_by(student_id=s.student_id, status='Present').count()
            absent_days = max(0, total_days_res - present_days)
            pct = round((present_days / total_days_res * 100), 1) if total_days_res > 0 else 0.0

            summary.append({
                'student_id': s.student_id,
                'name': s.name,
                'roll_number': s.roll_number,
                'course': s.course,
                'total_campus_days': total_days_res,
                'present_days': present_days,
                'absent_days': absent_days,
                'percentage': pct
            })

        return summary
